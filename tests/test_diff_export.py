"""
Teste unitário para validação da exportação dos relatórios de Diff (4 CSVs + 1 XLSX multi-abas).
"""

import os
import shutil
import pytest
from pathlib import Path
from src.c_aggregator.bridge import export_diff_to_files


def test_export_diff_to_files(tmp_path: Path):
    """
    Testa se export_diff_to_files gera os 4 arquivos CSV e a planilha Excel com 4 abas.
    """
    mock_stats = {
        "novos": [
            {
                "snowflake_id": 10001,
                "nome_instituicao": "Escola Nova Teste",
                "dependencia_adm": "Estadual",
                "endereco": "Rua das Flores, 100, Centro, Porto Alegre RS - 90000000",
                "municipio": "Porto Alegre",
                "uf": "RS"
            }
        ],
        "alterados": [
            {
                "id": 20002,
                "nome_instituicao": "Escola Alterada Teste",
                "dependencia_adm": "Federal",
                "diferencas_detectadas": ["endereco: Rua A -> Rua B"],
                "endereco_db": "Rua A, 10, Centro, Esteio RS",
                "endereco": "Rua B, 20, Centro, Esteio RS",
                "municipio": "Esteio",
                "uf": "RS"
            }
        ],
        "inativados": [
            {
                "id": 30003,
                "nome_instituicao": "Escola Inativada Teste",
                "dependencia_adm": "Privada",
                "endereco": "Av Brasil, 500, Dourados MS",
                "municipio": "Dourados",
                "uf": "MS"
            }
        ],
        "cursos_detalhes": [
            {
                "id": 10001,
                "nome_instituicao": "Escola Nova Teste",
                "tipo": "Nova Oferta",
                "curso": "Técnico em Informática"
            }
        ]
    }

    output_dir = tmp_path / "output_test"
    paths = export_diff_to_files(mock_stats, output_dir)

    # 1. Verifica se os 4 CSVs e o Excel existem
    assert Path(paths["csv_1"]).exists()
    assert Path(paths["csv_2"]).exists()
    assert Path(paths["csv_3"]).exists()
    assert Path(paths["csv_4"]).exists()
    assert Path(paths["excel"]).exists()

    # 2. Valida o conteúdo dos CSVs
    with open(paths["csv_1"], "r", encoding="utf-8") as f:
        content1 = f.read()
        assert "Escola Nova Teste" in content1
        assert "Porto Alegre" in content1

    with open(paths["csv_2"], "r", encoding="utf-8") as f:
        content2 = f.read()
        assert "Escola Alterada Teste" in content2
        assert "Rua A -> Rua B" in content2

    with open(paths["csv_3"], "r", encoding="utf-8") as f:
        content3 = f.read()
        assert "Escola Inativada Teste" in content3
        assert "Inativado (ativo=0)" in content3

    with open(paths["csv_4"], "r", encoding="utf-8") as f:
        content4 = f.read()
        assert "Técnico em Informática" in content4

    # 3. Valida a planilha Excel (.xlsx) e suas 4 abas
    try:
        import openpyxl
        wb = openpyxl.load_workbook(paths["excel"])
        sheet_names = wb.sheetnames
        assert "1. Instituições Novas" in sheet_names
        assert "2. Alterações Cadastrais" in sheet_names
        assert "3. Instituições Inativadas" in sheet_names
        assert "4. Detalhamento de Ofertas" in sheet_names

        ws1 = wb["1. Instituições Novas"]
        assert ws1.cell(row=2, column=2).value == "Escola Nova Teste"
    except ImportError:
        pass
